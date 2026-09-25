from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


def add_sheet(workbook, sheet_name, qna_list):

    worksheet = workbook.create_sheet(sheet_name)

    # Header
    worksheet.append([
        "Questions",
        "Answers"
    ])

    # Header formatting
    for cell in worksheet[1]:

        cell.font = Font(bold=True)

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # Data
    for item in qna_list:

        worksheet.append([
            item["question"],
            item["answer"]
        ])

    # Formatting
    worksheet.freeze_panes = "A2"

    worksheet.column_dimensions["A"].width = 55
    worksheet.column_dimensions["B"].width = 80

    for row in worksheet.iter_rows():

        for cell in row:

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )


def create_excel(
    english_qna,
    hindi_qna,
    marathi_qna,
    output_file
):

    workbook = Workbook()

    # Remove default sheet
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    add_sheet(
        workbook,
        "English",
        english_qna
    )

    add_sheet(
        workbook,
        "Hindi",
        hindi_qna
    )

    add_sheet(
        workbook,
        "Marathi",
        marathi_qna
    )

    workbook.save(output_file)